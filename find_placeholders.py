import docx
import re
import sys

doc_path = 'Book_with_Shipiki_Project_Documentation.docx'
doc = docx.Document(doc_path)

print("--- Searching for placeholders ---")
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if not text:
        continue
    # Look for common placeholder patterns like [Screenshot], [Figure], etc.
    if re.search(r'\[.*?(Screenshot|Figure|Insert).*?\]', text, re.IGNORECASE):
        print(f"[{i}] {text}")
    elif "screenshot" in text.lower() or "figure" in text.lower():
        print(f"[{i}] {text}")

print("--- Done ---")
