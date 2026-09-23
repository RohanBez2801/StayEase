import docx
from docx.shared import Inches
import os

doc_path = 'Book_with_Shipiki_Project_Documentation.docx'
out_path = 'Book_with_Shipiki_Project_Documentation_with_Screenshots.docx'
img_dir = r'C:\Users\rouxn\source\repos\StayEase\documentation\screenshots'

doc = docx.Document(doc_path)

# Map text snippets (to identify the paragraph) to the image filename to insert AFTER it
# We use lowercased snippets for robust matching
insert_map = {
    "14.1. system architecture": ['tech_doc_01_system_documentation.png'],
    "14.2. core database tables": ['tech_doc_02_database_schema.png'],
    "enter your destination city and travel dates": ['user_doc_01_homepage_search.png', 'user_doc_02_search_filled.png'],
    "browse the list of hotels and rooms shown": ['user_doc_03_search_results.png', 'user_doc_04_room_details.png'],
    "log in or create a free account": ['user_doc_05_login_page.png', 'user_doc_06_register_page.png'],
    "enter your payment details on the secure payment page": ['user_doc_07_checkout_page.png'],
    "a confirmation screen and email appear": ['user_doc_08_booking_confirmation.png'],
    "log in and open my bookings": ['user_doc_09_my_bookings.png'],
    "log in to the staff dashboard": ['user_doc_10_admin_dashboard.png'],
    "open rooms": ['user_doc_11_admin_rooms.png'], # Add room type
    "availability calendar": ['user_doc_12_admin_bookings.png'],
    "management reports": ['tech_doc_03_reports_analytics.png'],
    "admin functions only open to authorised staff": ['tech_doc_04_admin_guests.png']
}

print("Inserting images...")

inserted_count = 0
# We iterate by index because we might be inserting paragraphs
i = 0
while i < len(doc.paragraphs):
    para = doc.paragraphs[i]
    text = para.text.strip().lower()
    
    # Check if this paragraph matches any in our map
    for snippet, image_files in insert_map.items():
        if snippet in text:
            print(f"Found match for '{snippet}'")
            # We want to insert AFTER this paragraph. We do this by inserting BEFORE the next paragraph.
            # If it's the last paragraph, we just add it to the end.
            for img_file in image_files:
                img_path = os.path.join(img_dir, img_file)
                if os.path.exists(img_path):
                    if i + 1 < len(doc.paragraphs):
                        next_para = doc.paragraphs[i + 1]
                        new_p = next_para.insert_paragraph_before()
                        run = new_p.add_run()
                        run.add_picture(img_path, width=Inches(6.0))
                        
                        # Add a caption
                        caption_p = next_para.insert_paragraph_before(f"Screenshot: {img_file}")
                        caption_p.style = doc.styles['Caption'] if 'Caption' in doc.styles else doc.styles['Normal']
                        
                        i += 2 # Skip the newly inserted picture and caption paragraphs
                    else:
                        doc.add_picture(img_path, width=Inches(6.0))
                        doc.add_paragraph(f"Screenshot: {img_file}")
                    inserted_count += 1
                    print(f"  Inserted {img_file}")
                else:
                    print(f"  ERROR: Image not found - {img_path}")
            
            # Remove from map so we don't insert it again if the snippet appears twice
            del insert_map[snippet]
            break # Break out of inner loop so we don't modify dict while iterating
            
    i += 1

doc.save(out_path)
print(f"Successfully inserted {inserted_count} images. Saved to {out_path}")
